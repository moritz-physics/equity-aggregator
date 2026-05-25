"""Unit tests for the export layer (CSV / JSON / XLSX)."""

from __future__ import annotations

import json
from io import BytesIO

from openpyxl import load_workbook

from equity_aggregator.core.export import COLUMNS, to_csv, to_json, to_xlsx
from equity_aggregator.core.models import QueryResult


def _decoded_csv(result: QueryResult) -> list[list[str]]:
    text = to_csv(result).decode("utf-8-sig")
    import csv

    return list(csv.reader(text.splitlines()))


def test_to_csv_has_bom_and_header(sample_query_result: QueryResult) -> None:
    raw = to_csv(sample_query_result)
    assert raw.startswith(b"\xef\xbb\xbf"), "CSV must start with UTF-8 BOM"

    rows = _decoded_csv(sample_query_result)
    assert rows[0] == list(COLUMNS)


def test_to_csv_row_count_matches_constituents(
    sample_query_result: QueryResult,
) -> None:
    rows = _decoded_csv(sample_query_result)
    expected = sum(len(i.constituents) for i in sample_query_result.indices)
    assert len(rows) - 1 == expected  # minus header


def test_to_csv_formats_numbers_and_dates(sample_query_result: QueryResult) -> None:
    rows = _decoded_csv(sample_query_result)
    # ADS row is first.
    ads = rows[1]
    headers = list(COLUMNS)
    price = ads[headers.index("Price")]
    yield_ = ads[headers.index("TTM Div Yield %")]
    market_cap = ads[headers.index("Market Cap")]
    ex_div = ads[headers.index("Ex-Div Date")]

    assert price == "154.50"
    assert yield_ == "1.81"
    assert market_cap == "27,400,000,000"
    assert ex_div == "2026-05-08"


def test_to_csv_empty_cells_for_none(sample_query_result: QueryResult) -> None:
    rows = _decoded_csv(sample_query_result)
    headers = list(COLUMNS)
    # XYZ row has no financials.
    xyz = next(r for r in rows[1:] if r[0] == "XYZ.DE")
    assert xyz[headers.index("Price")] == ""
    assert xyz[headers.index("ISIN")] == ""
    assert xyz[headers.index("Sector")] == ""


def test_to_json_is_valid_with_indices_key(
    sample_query_result: QueryResult,
) -> None:
    payload = to_json(sample_query_result)
    parsed = json.loads(payload)
    assert isinstance(parsed, dict)
    assert "indices" in parsed
    assert parsed["total_constituents"] == 3
    assert parsed["indices"][0]["name"] == "DAX"
    assert parsed["indices"][0]["constituents"][0]["ticker"] == "ADS.DE"
    # Dates serialise as ISO strings, not epoch.
    assert parsed["indices"][0]["constituents"][0]["ex_div_date"] == "2026-05-08"


def test_to_json_is_indented(sample_query_result: QueryResult) -> None:
    payload = to_json(sample_query_result)
    # Pretty-printed JSON has newlines and leading spaces.
    assert "\n" in payload
    assert "  " in payload


def test_to_xlsx_bytes_non_empty(sample_query_result: QueryResult) -> None:
    raw = to_xlsx(sample_query_result)
    assert isinstance(raw, bytes)
    assert len(raw) > 1000


def test_to_xlsx_sheet_name_and_layout(
    sample_query_result: QueryResult,
) -> None:
    raw = to_xlsx(sample_query_result)
    wb = load_workbook(BytesIO(raw))
    assert wb.sheetnames == ["DAX"]

    ws = wb["DAX"]
    # Row 1 is the title row — contains the index name.
    assert "DAX" in str(ws.cell(row=1, column=1).value)

    # Row 2 is the column header row (all caps).
    header_row = [ws.cell(row=2, column=c).value for c in range(1, 12)]
    assert header_row[0] == "TICKER"
    assert header_row[1] == "COMPANY"
    assert all(isinstance(v, str) and v == v.upper() for v in header_row)

    # 3 constituents → rows 3..5; ws.max_row should be 5.
    assert ws.max_row == 2 + 3
    assert ws.cell(row=3, column=1).value == "ADS.DE"
    assert ws.cell(row=4, column=1).value == "SAP.DE"
    assert ws.cell(row=5, column=1).value == "XYZ.DE"


def test_to_xlsx_title_row_is_merged(sample_query_result: QueryResult) -> None:
    raw = to_xlsx(sample_query_result)
    wb = load_workbook(BytesIO(raw))
    ws = wb["DAX"]
    merged_ranges = [str(r) for r in ws.merged_cells.ranges]
    assert any(r.startswith("A1:") and r.endswith("1") for r in merged_ranges)


def test_to_xlsx_no_none_strings(sample_query_result: QueryResult) -> None:
    raw = to_xlsx(sample_query_result)
    wb = load_workbook(BytesIO(raw))
    ws = wb["DAX"]
    for row in ws.iter_rows(min_row=3, values_only=True):
        for cell in row:
            assert cell != "None"
            assert cell is None or not (
                isinstance(cell, str) and cell.lower() == "none"
            )


def test_to_xlsx_has_freeze_panes(sample_query_result: QueryResult) -> None:
    raw = to_xlsx(sample_query_result)
    wb = load_workbook(BytesIO(raw))
    ws = wb["DAX"]
    assert ws.freeze_panes == "A3"
