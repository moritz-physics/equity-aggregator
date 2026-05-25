"""Export ``QueryResult`` snapshots to CSV, JSON, and XLSX.

CSV is UTF-8 with BOM (so Excel opens it cleanly). XLSX uses a professional
finance-terminal style — navy title and header bands with alternating
white/light-grey data rows. JSON is pretty-printed via Pydantic so dates
serialise as ISO strings.
"""

from __future__ import annotations

import csv
import io
from datetime import UTC, date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook as WorkbookT

from equity_aggregator.core.models import Constituent, Index, QueryResult

COLUMNS: tuple[str, ...] = (
    "Ticker",
    "Company",
    "ISIN",
    "ISIN Source",
    "Country",
    "Index Weight %",
    "Price",
    "Currency",
    "Ex-Div Date",
    "TTM Div Yield %",
    "Beta",
    "Market Cap",
    "Sector",
)

# ── XLSX palette (Excel-safe, finance-terminal style) ────────────────────
HEADER_FILL = "1F3864"   # deep navy
HEADER_FONT = "FFFFFF"   # white
TITLE_FILL = "0A1628"    # near-black navy for title row
TITLE_FONT = "FFFFFF"
ROW_ALT_FILL = "F2F2F2"  # light grey alternate rows
ROW_FILL = "FFFFFF"      # white base rows
BORDER_COLOR = "BFBFBF"  # light grey border
ACCENT = "2563EB"        # blue accent for header underline

# (name, width, kind, number_format)
XLSX_COLUMNS: tuple[tuple[str, int, str, str], ...] = (
    ("Ticker",        12, "text",   "@"),
    ("Company",       32, "text",   "@"),
    ("ISIN",          16, "text",   "@"),
    ("Country",        8, "text",   "@"),
    ("Weight %",      11, "number", '0.0000"%"'),
    ("Price",         12, "number", "#,##0.00"),
    ("Currency",       8, "text",   "@"),
    ("Ex-Div Date",   14, "date",   "YYYY-MM-DD"),
    ("TTM Div Yield", 14, "number", '0.00"%"'),
    ("Beta",          10, "number", "0.00"),
    ("Market Cap",    18, "number", "#,##0"),
    ("Sector",        22, "text",   "@"),
)


def _csv_cell(value: Any, *, digits: int = 2) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.{digits}f}"
    return str(value)


def _csv_row(c: Constituent) -> list[str]:
    return [
        c.ticker,
        c.name,
        c.isin or "",
        c.isin_source or "",
        c.country,
        _csv_cell(c.weight, digits=4),
        _csv_cell(c.price),
        c.currency or "",
        c.ex_div_date.isoformat() if c.ex_div_date is not None else "",
        _csv_cell(c.ttm_div_yield),
        _csv_cell(c.beta, digits=3),
        f"{c.market_cap:,.0f}" if c.market_cap is not None else "",
        c.sector or "",
    ]


def to_csv(result: QueryResult) -> bytes:
    """UTF-8 BOM CSV. Excel auto-detects encoding when the BOM is present."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(COLUMNS)
    for index in result.indices:
        for constituent in index.constituents:
            writer.writerow(_csv_row(constituent))
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def to_json(result: QueryResult) -> str:
    """Pretty-printed JSON via Pydantic so dates serialise as ISO strings."""
    return result.model_dump_json(indent=2)


def _xlsx_value(c: Constituent, name: str) -> Any:
    """Return native value for an XLSX cell — empty string for missing."""
    if name == "Ticker":
        return c.ticker
    if name == "Company":
        return c.name
    if name == "ISIN":
        return c.isin or ""
    if name == "Country":
        return c.country
    if name == "Weight %":
        return c.weight if c.weight is not None else ""
    if name == "Price":
        return c.price if c.price is not None else ""
    if name == "Currency":
        return c.currency or ""
    if name == "Ex-Div Date":
        return c.ex_div_date if c.ex_div_date is not None else ""
    if name == "TTM Div Yield":
        return c.ttm_div_yield if c.ttm_div_yield is not None else ""
    if name == "Beta":
        return c.beta if c.beta is not None else ""
    if name == "Market Cap":
        return c.market_cap if c.market_cap is not None else ""
    if name == "Sector":
        return c.sector or ""
    return ""


def _style_xlsx_sheet(ws: Any, index: Index) -> None:
    n_cols = len(XLSX_COLUMNS)

    title_fill = PatternFill("solid", fgColor=TITLE_FILL)
    title_font = Font(bold=True, color=TITLE_FONT, size=11)
    title_align = Alignment(horizontal="left", vertical="center")

    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color=HEADER_FONT, size=10)
    header_align = Alignment(horizontal="center", vertical="center")
    accent_side = Side(border_style="thin", color=ACCENT)
    header_border = Border(bottom=accent_side)

    row_fill = PatternFill("solid", fgColor=ROW_FILL)
    alt_fill = PatternFill("solid", fgColor=ROW_ALT_FILL)
    grid_side = Side(border_style="thin", color=BORDER_COLOR)
    grid_border = Border(
        left=grid_side, right=grid_side, top=grid_side, bottom=grid_side
    )
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    center_align = Alignment(horizontal="center", vertical="center")

    # ── Row 1: title (merged) ────────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title_text = (
        f"{index.name} — Equity Constituents  |  "
        f"{date.today().isoformat()}  |  "
        f"Source: Yahoo Finance (unofficial)"
    )
    for col_idx in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = title_fill
        cell.font = title_font
        cell.alignment = title_align
    ws.cell(row=1, column=1).value = title_text
    ws.row_dimensions[1].height = 22

    # ── Row 2: column headers (all caps) ─────────────────────────────────
    for col_idx, (name, _w, _kind, _fmt) in enumerate(XLSX_COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=name.upper())
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = header_border
    ws.row_dimensions[2].height = 18

    # ── Rows 3+: data ────────────────────────────────────────────────────
    for row_offset, constituent in enumerate(index.constituents):
        excel_row = row_offset + 3
        fill = alt_fill if row_offset % 2 else row_fill
        for col_idx, (name, _w, kind, fmt) in enumerate(XLSX_COLUMNS, start=1):
            value = _xlsx_value(constituent, name)
            cell = ws.cell(row=excel_row, column=col_idx, value=value)
            cell.fill = fill
            cell.border = grid_border
            cell.number_format = fmt
            if kind == "number":
                cell.alignment = right_align
            elif kind == "date":
                cell.alignment = center_align
            else:
                cell.alignment = left_align
        ws.row_dimensions[excel_row].height = 15

    # ── Column widths ────────────────────────────────────────────────────
    for col_idx, (_n, width, _k, _f) in enumerate(XLSX_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ── Freeze title + header ────────────────────────────────────────────
    ws.freeze_panes = "A3"


def to_xlsx(result: QueryResult) -> bytes:
    """Themed multi-sheet XLSX — one sheet per index, freeze panes at A3."""
    wb: WorkbookT = Workbook()
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)

    if not result.indices:
        ws = wb.create_sheet(title="Empty")
        empty_index = Index(
            name="Empty",
            country="",
            constituents=[],
            fetched_at=datetime.now(tz=UTC),
            source="empty",
        )
        _style_xlsx_sheet(ws, empty_index)
    else:
        for index in result.indices:
            ws = wb.create_sheet(title=index.name[:31])
            _style_xlsx_sheet(ws, index)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
